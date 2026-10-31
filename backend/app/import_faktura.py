"""Загрузка тестовых проектов кейса из Keys_YuFU_Faktura.xlsx (лист «Список проектов»).

Запуск: python -m app.import_faktura путь/к/Keys_YuFU_Faktura.xlsx [--author admin@example.com]

Повторный запуск не создаёт дублей: проект с тем же клиентом и названием пропускается.
Вероятность берётся из справочника этапов, а не из столбца «%» файла: по ТЗ она зависит от этапа.
Месячные суммы файла записываются строками выручки в тысячах рублей со статусом «прогнозное начисление».
"""

import argparse
import re
import secrets
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.db import SessionLocal
from app.models import (
    AuditLog,
    Client,
    Project,
    ProjectFinancialEntry,
    ReferenceValue,
    Role,
    Segment,
    Service,
    User,
)
from app.reference_data import SPEC_SERVICES
from app.services.project_history import project_snapshot

SHEET = "Список проектов"
HEADER_ROW = 3
MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь",
          "декабрь"]
UNKNOWN_SEGMENT = "Сегмент не указан"
TRANSLIT = str.maketrans({
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i", "й": "y",
    "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f",
    "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
})


@dataclass
class ImportResult:
    created: int = 0
    skipped: int = 0
    clients: int = 0
    managers: int = 0
    warnings: list[str] = field(default_factory=list)


def norm(value) -> str:
    """Убирает неразрывные пробелы и повторные пробелы: «3.\\xa0КП» → «3. КП»."""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip() if value is not None else ""


def read_rows(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True, read_only=True)[SHEET]
    rows = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = [norm(h) for h in next(rows)]
    result = []
    for values in rows:
        row = dict(zip(header, values, strict=False))
        if row.get("Клиент") and row.get("Описание проекта"):
            result.append(row)
    return result


def _reference(db: Session, category: str) -> dict[str, ReferenceValue]:
    refs = db.scalars(select(ReferenceValue).where(ReferenceValue.category == category)).all()
    return {norm(r.name).casefold(): r for r in refs}


def _match(refs: dict[str, ReferenceValue], value) -> ReferenceValue | None:
    key = norm(value).casefold()
    if not key:
        return None
    if key in refs:
        return refs[key]
    # «Интеграционные проекты-ЕГЭ» → «Интеграционные проекты»
    return next((ref for name, ref in refs.items() if key.startswith(name)), None)


def _manager(db: Session, surname: str, cache: dict[str, User], result: ImportResult) -> User:
    if surname in cache:
        return cache[surname]
    user = db.scalar(select(User).where(User.full_name == surname))
    if user is None:
        login = surname.casefold().translate(TRANSLIT) or "manager"
        email = f"{login}@faktura.example.com"
        n = 2
        while db.scalar(select(User.id).where(User.email == email)) is not None:
            email = f"{login}{n}@faktura.example.com"
            n += 1
        # Пароль случайный: доступ менеджеру выдаёт администратор в разделе «Пользователи».
        password_hash = hash_password(secrets.token_urlsafe(16))
        user = User(email=email, full_name=surname, role=Role.user, password_hash=password_hash)
        db.add(user)
        db.flush()
        result.managers += 1
    cache[surname] = user
    return user


def _created_at(value) -> datetime:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    return datetime.now(UTC)


def import_rows(db: Session, rows: list[dict], author: User) -> ImportResult:
    result = ImportResult()
    stages = _reference(db, "customer_stage")
    payment_types = _reference(db, "payment_type")
    assessments = _reference(db, "assessment_status")
    revenue_status = db.scalar(select(ReferenceValue).where(
        ReferenceValue.category == "revenue_status", ReferenceValue.code == "forecast",
    ))
    segments = {s.name: s for s in db.scalars(select(Segment))}
    services = {s.name: s for s in db.scalars(select(Service))}
    managers: dict[str, User] = {}

    for number, row in enumerate(rows, start=HEADER_ROW + 1):
        client_name, project_name = norm(row["Клиент"]), norm(row["Описание проекта"])
        stage = _match(stages, row.get("Этап"))
        if stage is None:
            result.warnings.append(f"Строка {number}: этап «{norm(row.get('Этап'))}» не найден, строка пропущена")
            result.skipped += 1
            continue

        segment_name = norm(row.get("Сегмент"))
        if segment_name not in segments:
            if segment_name:
                result.warnings.append(f"Строка {number}: сегмент «{segment_name}» не из справочника")
            segment_name = UNKNOWN_SEGMENT
            if segment_name not in segments:
                segments[segment_name] = Segment(name=segment_name)
                db.add(segments[segment_name])
                db.flush()

        service_name = norm(row.get("Услуга крупно")) or "Услуга не указана"
        if service_name not in services:
            services[service_name] = Service(name=service_name, group_name=SPEC_SERVICES.get(service_name))
            db.add(services[service_name])
            db.flush()

        inn = re.sub(r"\D", "", norm(row.get("Столбец1")))[:12] or None
        client = db.scalar(select(Client).where(Client.inn == inn)) if inn else None
        if client is None:
            client = db.scalar(select(Client).where(Client.name == client_name))
        if client is None:
            client = Client(name=client_name, inn=inn or f"F{number:011d}", segment=segments[segment_name])
            db.add(client)
            db.flush()
            result.clients += 1

        if db.scalar(select(Project.id).where(Project.client_id == client.id, Project.name == project_name)):
            result.skipped += 1
            continue

        created_at = _created_at(row.get("Дата"))
        year = int(row.get("Год реализации") or created_at.year)
        months = [Decimal(str(row.get(m) or 0)) for m in MONTHS]
        total = sum(months, Decimal("0"))
        industry = norm(row.get("Отраслевое решени")).casefold() == "да"
        assessment = _match(assessments, row.get("Принимаемый к оценке"))
        payment_type = _match(payment_types, row.get("Тип платежа"))
        start = created_at.date()

        project = Project(
            name=project_name,
            description="",
            client_id=client.id,
            segment_id=segments[segment_name].id,
            manager_id=_manager(db, norm(row.get("Менеджер")) or "Менеджер не указан", managers, result).id,
            service_id=services[service_name].id,
            stage_scheme="customer",
            customer_stage_id=stage.id,
            probability=stage.probability or 0,
            amount=(total * 1000).quantize(Decimal("0.01")),
            cost=Decimal("0"),
            start_date=start,
            expected_close_date=max(date(year, 12, 31), start),
            payment_type_id=payment_type.id if payment_type else None,
            accepted_for_forecast=assessment is not None,
            assessment_status_id=assessment.id if assessment else None,
            is_industry_solution=industry,
            industry_manager=(norm(row.get("Отраслевой менеджер")) or None) if industry else None,
            industry_project_number=(norm(row.get("№ проекта ПО")) or None) if industry else None,
            implementation_year=year,
            project_status=norm(row.get("Статус"))[:1000] or None,
            realized_work=norm(row.get("Что сделано за 14 дней"))[:1000],
            period_plan=norm(row.get("План на 2 недели"))[:1000],
            created_at=created_at,
            updated_at=created_at,
        )
        project.financial_entries = [
            ProjectFinancialEntry(kind="revenue", year=year, month=i + 1, amount=amount, vat_included=False,
                                  status=revenue_status)
            for i, amount in enumerate(months)
            if amount > 0
        ]
        db.add(project)
        db.flush()
        db.add(AuditLog(
            user_id=author.id, action="create", entity="project", entity_id=project.id, created_at=created_at,
            payload={"name": project.name, "after": project_snapshot(db, project), "author_name": author.full_name,
                     "source": "Keys_YuFU_Faktura.xlsx"},
        ))
        result.created += 1
    db.commit()
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("path", type=Path)
    parser.add_argument("--author", default="admin@example.com", help="от чьего имени записать создание в историю")
    args = parser.parse_args()
    with SessionLocal() as db:
        author = db.scalar(select(User).where(User.email == args.author))
        if author is None:
            raise SystemExit(f"Пользователь {args.author} не найден; сначала выполните python -m app.seed")
        result = import_rows(db, read_rows(args.path), author)
    for warning in result.warnings:
        print("Предупреждение:", warning)
    print(f"Создано проектов: {result.created}, пропущено: {result.skipped}, "
          f"новых клиентов: {result.clients}, новых менеджеров: {result.managers}")


if __name__ == "__main__":
    main()
