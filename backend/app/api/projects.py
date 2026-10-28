from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.core import cache
from app.db import get_db
from app.deps import PROJECT_EDITORS, get_current_user, manager_scope, require_roles
from app.history_schemas import ProjectHistoryPage
from app.models import (
    CLOSED_STAGES,
    STAGE_LABELS,
    AuditLog,
    Client,
    Project,
    ProjectFinancialEntry,
    ReferenceValue,
    Role,
    Service,
    Stage,
    User,
)
from app.project_write_schemas import ProjectSaveIn
from app.schemas import (
    LEGACY_STAGE_SCOPE_NOTE,
    NameCheckOut,
    ProjectFinancialEntryIn,
    ProjectFinancialEntryOut,
    ProjectIn,
    ProjectOut,
    ProjectPage,
)
from app.services import audit
from app.services.project_field_label_export import PROJECT_EXPORT_FIELDS, write_project_export_headers
from app.services.project_field_labels import (
    effective_project_field_labels,
    same_project_values,
    with_snapshot_field_labels,
)
from app.services.project_history import financial_snapshot, history_items, project_snapshot, same_finances

router = APIRouter(prefix="/projects", tags=["projects"])

SORT_FIELDS = {
    "name": Project.name,
    "client_inn": Client.inn,
    "implementation_year": Project.implementation_year,
    "customer_stage_id": Project.customer_stage_id,
    "amount": Project.amount,
    "start_date": Project.start_date,
    "expected_close_date": Project.expected_close_date,
    "updated_at": Project.updated_at,
}
SORT_PATTERN = (
    r"^-?(name|client_inn|implementation_year|customer_stage_id|"
    r"amount|start_date|expected_close_date|updated_at)$"
)
EXPORT_LIMIT = 10_000


def _xlsx_timestamp(value: datetime) -> str:
    # Excel не поддерживает tzinfo: явный ISO 8601 в UTC сохраняет смысл времени.
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).isoformat()


def to_out(p: Project) -> dict:
    margin = p.amount - p.cost
    return {
        "id": p.id,
        "name": p.name,
        "description": p.description,
        "client": p.client,
        "client_inn": p.client.inn,
        "segment": p.client.segment,
        "manager": p.manager,
        "service": p.service,
        "stage": p.stage,
        "stage_scheme": p.stage_scheme,
        "amount": p.amount,
        "cost": p.cost,
        "margin": margin,
        "margin_pct": round(float(margin / p.amount * 100), 1) if p.amount else None,
        "probability": p.probability,
        "start_date": p.start_date,
        "expected_close_date": p.expected_close_date,
        "closed_at": p.closed_at,
        "payment_type": p.payment_type,
        "customer_stage": p.customer_stage,
        "indicator": p.indicator,
        "accepted_for_forecast": p.accepted_for_forecast,
        "assessment_status": p.assessment_status,
        "subsidiary_implementation": p.subsidiary_implementation,
        "management_control": p.management_control,
        "is_industry_solution": p.is_industry_solution,
        "industry_manager": p.industry_manager,
        "industry_project_number": p.industry_project_number,
        "implementation_year": p.implementation_year,
        "project_status": p.project_status,
        "realized_work": p.realized_work,
        "period_plan": p.period_plan,
        "created_at": p.created_at,
        "updated_at": p.updated_at,
        "archived_at": p.archived_at,
        "financial_entries": p.financial_entries,
    }


def _filtered_query(
    user: User,
    q: str,
    segment_id: list[int],
    service_id: list[int],
    manager_id: list[int],
    stage: list[Stage],
    open_only: bool,
    archived: bool = False,
    customer_stage_id: int | None = None,
    implementation_year: int | None = None,
) -> Select:
    stmt = select(Project).options(selectinload(Project.financial_entries)).join(Client, Client.id == Project.client_id)
    if archived and user.role != Role.admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Архив доступен только администратору")
    stmt = stmt.where(Project.archived_at.is_not(None) if archived else Project.archived_at.is_(None))
    scope = manager_scope(user)
    if scope is not None:
        stmt = stmt.where(Project.manager_id == scope)
    q = q.strip()
    if q:
        stmt = stmt.where(or_(Project.name.ilike(f"%{q}%"), Client.name.ilike(f"%{q}%"), Client.inn.startswith(q)))
    if segment_id:
        stmt = stmt.where(Client.segment_id.in_(segment_id))
    if service_id:
        stmt = stmt.where(Project.service_id.in_(service_id))
    if manager_id:
        stmt = stmt.where(Project.manager_id.in_(manager_id))
    if stage:
        stmt = stmt.where(Project.stage.in_(stage))
    if customer_stage_id is not None:
        stmt = stmt.where(Project.customer_stage_id == customer_stage_id)
    if implementation_year is not None:
        stmt = stmt.where(Project.implementation_year == implementation_year)
    if open_only:
        stmt = stmt.where(Project.stage_scheme == "legacy", Project.stage.notin_(CLOSED_STAGES))
    return stmt


class ListParams:
    def __init__(
        self,
        q: str = Query(default="", max_length=100),
        segment_id: list[int] = Query(default=[]),
        service_id: list[int] = Query(default=[]),
        manager_id: list[int] = Query(default=[]),
        stage: list[Stage] = Query(default=[]),
        open_only: bool = False,
        archived: bool = False,
        sort: str = Query(default="-updated_at", pattern=SORT_PATTERN),
        customer_stage_id: int | None = Query(default=None, ge=1),
        implementation_year: int | None = Query(default=None, ge=1900, le=2200),
        date_from: date | None = None,
        date_to: date | None = None,
        period_preset: Literal["week", "month", "quarter", "custom"] | None = None,
    ):
        if (date_from is None) != (date_to is None) or (period_preset is not None and date_from is None):
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Укажите начало и конец периода")
        if date_from is not None and date_to is not None:
            if date_from > date_to:
                raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Начало периода не может быть позже конца")
            if (date_to - date_from).days > 36600:
                raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Период не может превышать 36600 дней")
        self.q, self.segment_id, self.service_id = q, segment_id, service_id
        self.manager_id, self.stage, self.open_only, self.sort = manager_id, stage, open_only, sort
        self.archived = archived
        self.customer_stage_id, self.implementation_year = customer_stage_id, implementation_year

    def query(self, user: User) -> Select:
        stmt = _filtered_query(
            user, self.q, self.segment_id, self.service_id, self.manager_id, self.stage, self.open_only,
            archived=self.archived, customer_stage_id=self.customer_stage_id,
            implementation_year=self.implementation_year,
        )
        field = self.sort.lstrip("-")
        column = SORT_FIELDS[field]
        order = column.desc() if self.sort.startswith("-") else column.asc()
        if field in {"customer_stage_id", "implementation_year"}:
            order = order.nulls_last()
        return stmt.order_by(order, Project.id.desc())


@router.get("", response_model=ProjectPage)
def list_projects(
    params: ListParams = Depends(),
    page: int = Query(default=1, ge=1),
    size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    stmt = params.query(user)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery()))
    items = db.scalars(stmt.offset((page - 1) * size).limit(size)).unique().all()
    return {
        "items": [to_out(p) for p in items], "total": total, "page": page, "size": size,
        "stage_scope": "legacy" if params.open_only else "all",
        "scope_note": LEGACY_STAGE_SCOPE_NOTE if params.open_only else None,
    }


@router.get("/check-name", response_model=NameCheckOut)
def check_name(
    name: str = Query(min_length=1, max_length=300),
    client_id: int = Query(),
    exclude_id: int | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> dict:
    """Проверка на дубль во время ввода (до отправки формы)."""
    return {"available": _find_duplicate(db, " ".join(name.split()), client_id, exclude_id) is None}


@router.get("/export.xlsx")
def export_projects(
    params: ListParams = Depends(),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(Role.analyst, Role.admin)),
) -> StreamingResponse:
    projects = db.scalars(params.query(user).limit(EXPORT_LIMIT)).unique().all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Проекты (6 стадий)" if params.open_only else "Проекты"
    if params.open_only:
        wb.create_sheet("Область данных").append([LEGACY_STAGE_SCOPE_NOTE])
        wb.properties.description = LEGACY_STAGE_SCOPE_NOTE
    write_project_export_headers(ws, effective_project_field_labels(db))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for p in projects:
        ws.append(
            [
                p.id,
                p.name,
                p.client.name,
                p.client.inn,
                p.client.segment.name,
                p.service.name,
                p.manager.full_name,
                STAGE_LABELS[p.stage] if p.stage is not None else "Не применяется",
                float(p.amount),
                float(p.cost),
                float(p.amount - p.cost),
                p.probability,
                p.start_date,
                p.expected_close_date,
                p.closed_at,
                p.payment_type.name if p.payment_type else None,
                p.customer_stage.name if p.customer_stage else None,
                p.indicator.name if p.indicator else None,
                p.accepted_for_forecast,
                p.assessment_status.name if p.assessment_status else None,
                p.subsidiary_implementation,
                p.management_control,
                p.is_industry_solution,
                p.industry_manager,
                p.industry_project_number,
                p.implementation_year,
                p.project_status,
                p.realized_work,
                p.period_plan,
                _xlsx_timestamp(p.created_at),
            ]
        )
        # Пользовательский текст должен оставаться текстом, даже если начинается с '='.
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    widths = [7, 40, 35, 14, 22, 26, 24, 16, 16, 16, 16, 12, 12, 14, 12] + [22] * (len(PROJECT_EXPORT_FIELDS) - 15)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="projects_{date.today():%Y%m%d}.xlsx"'},
    )


def _get_visible(db: Session, project_id: int, user: User, *, for_update: bool = False) -> Project:
    query = select(Project).where(Project.id == project_id)
    if for_update:
        query = query.with_for_update(of=Project).execution_options(populate_existing=True)
    project = db.scalar(query)
    scope = manager_scope(user)
    if (project is None or (scope is not None and project.manager_id != scope)
            or (project.archived_at is not None and user.role != Role.admin)):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Проект не найден")
    return project


def _require_active(project: Project) -> None:
    if project.archived_at is not None:
        raise HTTPException(status.HTTP_409_CONFLICT, "Сначала восстановите проект из архива")


def _set_archived(db: Session, project_id: int, user: User, *, archived: bool) -> dict:
    project = _get_visible(db, project_id, user, for_update=True)
    if (project.archived_at is not None) == archived:
        return to_out(project)
    before = project.archived_at.isoformat() if project.archived_at else None
    field_labels = effective_project_field_labels(db)
    project.archived_at = datetime.now(UTC) if archived else None
    project.archived_by_id = user.id if archived else None
    project.updated_at = func.now()
    audit.record(
        db, user.id, "archive" if archived else "restore", "project", project.id,
        before=with_snapshot_field_labels({"archived_at": before, "_stage_scheme": project.stage_scheme}, field_labels),
        after=with_snapshot_field_labels(
            {"archived_at": project.archived_at.isoformat() if project.archived_at else None,
             "_stage_scheme": project.stage_scheme}, field_labels,
        ),
        author_name=user.full_name,
    )
    db.commit()
    cache.bump("analytics")
    db.refresh(project)
    return to_out(project)


@router.post("/{project_id}/archive", response_model=ProjectOut)
def archive_project(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(Role.admin)),
) -> dict:
    return _set_archived(db, project_id, user, archived=True)


@router.post("/{project_id}/restore", response_model=ProjectOut)
def restore_project(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(Role.admin)),
) -> dict:
    return _set_archived(db, project_id, user, archived=False)


@router.get("/{project_id}", response_model=ProjectOut)
def get_project(project_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return to_out(_get_visible(db, project_id, user))


@router.get("/{project_id}/history", response_model=ProjectHistoryPage)
def get_project_history(
    project_id: int,
    limit: int = Query(default=10, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    _get_visible(db, project_id, user)
    rows = db.execute(
        select(AuditLog, User.full_name)
        .outerjoin(User, User.id == AuditLog.user_id)
        .where(AuditLog.entity == "project", AuditLog.entity_id == project_id)
        .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        .offset(offset)
        .limit(limit + 1)
    ).all()
    return {"items": history_items(db, rows[:limit]), "limit": limit, "offset": offset, "has_more": len(rows) > limit}


def _find_duplicate(db: Session, name: str, client_id: int, exclude_id: int | None) -> Project | None:
    # Сравнение без учёта регистра в Python: одинаково работает с кириллицей в любой СУБД.
    # Проектов у одного клиента немного, поэтому выборка дешёвая.
    stmt = select(Project).where(Project.client_id == client_id)
    if exclude_id is not None:
        stmt = stmt.where(Project.id != exclude_id)
    key = name.casefold()
    return next((p for p in db.scalars(stmt).unique() if p.name.casefold() == key), None)


def _check_integrity_conflict(db: Session, data: ProjectIn, project_id: int | None = None) -> None:
    db.rollback()
    if _find_duplicate(db, data.name, data.client_id, project_id) is not None:
        raise HTTPException(
            status.HTTP_409_CONFLICT, {"field": "name", "message": "У этого клиента уже есть проект с таким названием"}
        ) from None


def _apply(
    db: Session, project: Project, data: ProjectIn, user: User, locked_refs: dict[int, ReferenceValue],
) -> None:
    customer = project.stage_scheme == "customer"
    stage_changed = customer and (project.id is None or data.customer_stage_id != project.customer_stage_id)
    if customer:
        if data.stage is not None:
            raise HTTPException(422, {
                "field": "stage", "message": "Для новых проектов выберите этап по ТЗ; прежняя стадия не применяется",
            })
        if data.customer_stage_id is None:
            raise HTTPException(422, {"field": "customer_stage_id", "message": "Выберите этап по ТЗ"})
    else:
        if "customer_stage_id" in data.model_fields_set and data.customer_stage_id != project.customer_stage_id:
            raise HTTPException(422, {
                "field": "customer_stage_id", "message": "Вспомогательный этап прежнего проекта изменять нельзя",
            })
        if "stage" in data.model_fields_set and data.stage is None:
            raise HTTPException(422, {"field": "stage", "message": "Укажите прежнюю стадию проекта"})
        if "probability" in data.model_fields_set and data.probability is None:
            raise HTTPException(422, {"field": "probability", "message": "Укажите вероятность проекта"})
    if db.get(Client, data.client_id) is None:
        raise HTTPException(422, {"field": "client_id", "message": "Клиент не найден"})
    service = db.scalar(
        select(Service).where(Service.id == data.service_id).with_for_update(read=True)
        .execution_options(populate_existing=True)
    )
    retained_service = project.id is not None and data.service_id == project.service_id
    if service is None or (not service.is_active and not retained_service):
        raise HTTPException(422, {"field": "service_id", "message": "Услуга не найдена"})

    expected_categories = {
        "payment_type_id": "payment_type",
        "customer_stage_id": "customer_stage",
        "indicator_id": "project_indicator",
        "assessment_status_id": "assessment_status",
    }
    references: dict[str, ReferenceValue] = {}
    for field, category in expected_categories.items():
        if field == "customer_stage_id" and not customer:
            continue
        reference_id = getattr(data, field)
        if reference_id is None:
            continue
        reference = locked_refs.get(reference_id)
        retained = project.id is not None and reference_id == getattr(project, field)
        if reference is None or (not reference.is_active and not retained) or reference.category != category:
            raise HTTPException(422, {"field": field, "message": "Значение справочника не найдено"})
        references[field] = reference

    if user.role == Role.user:
        manager_id = user.id
    else:
        manager_id = data.manager_id or user.id
        manager = db.get(User, manager_id)
        if manager is None or not manager.is_active:
            raise HTTPException(422, {"field": "manager_id", "message": "Ответственный не найден"})

    if _find_duplicate(db, data.name, data.client_id, project.id):
        raise HTTPException(
            status.HTTP_409_CONFLICT, {"field": "name", "message": "У этого клиента уже есть проект с таким названием"}
        )

    customer_stage = references.get("customer_stage_id")
    if stage_changed and customer_stage is not None and customer_stage.probability is None:
        raise HTTPException(422, {"field": "customer_stage_id", "message": "У этапа не задана вероятность"})

    was_closed = project.stage in CLOSED_STAGES if project.stage else False
    for field, value in data.model_dump(
        exclude={"manager_id", "financial_entries", "stage", "customer_stage_id", "probability"},
    ).items():
        setattr(project, field, value)
    project.manager_id = manager_id
    if customer:
        project.stage = None
        project.customer_stage_id = data.customer_stage_id
        if stage_changed:
            project.probability = customer_stage.probability
        project.closed_at = None
        return
    if data.stage is not None:
        project.stage = data.stage
    if data.probability is not None:
        project.probability = data.probability
    if project.stage in CLOSED_STAGES:
        if not was_closed or project.closed_at is None:
            project.closed_at = max(date.today(), data.start_date)
    else:
        project.closed_at = None


@router.post("", response_model=ProjectOut, status_code=status.HTTP_201_CREATED)
def create_project(
    data: ProjectSaveIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*PROJECT_EDITORS)),
) -> dict:
    project = Project(cost=Decimal("0"), stage_scheme="customer")
    try:
        locked_refs = _lock_references(db, data, data.financial_entries or [])
        references = _validate_financial_entries(data.financial_entries or [], locked_refs)
        _apply(db, project, data, user, locked_refs)
        db.add(project)
        db.flush()
        audit.record(
            db, user.id, "create", "project", project.id,
            name=project.name, after=project_snapshot(db, project), author_name=user.full_name,
        )
        if data.financial_entries is not None:
            _replace_financial_entries(db, project, data.financial_entries, references, user)
        db.commit()
    except IntegrityError:
        _check_integrity_conflict(db, data)
        raise
    except Exception:
        db.rollback()
        raise
    cache.bump("analytics")
    db.refresh(project)
    return to_out(project)


@router.put("/{project_id}", response_model=ProjectOut)
def update_project(
    project_id: int,
    data: ProjectSaveIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*PROJECT_EDITORS)),
) -> dict:
    project = _get_visible(db, project_id, user, for_update=True)
    _require_active(project)
    try:
        locked_refs = _lock_references(db, data, data.financial_entries or [])
        references = _validate_financial_entries(data.financial_entries or [], locked_refs, project)
        field_labels = effective_project_field_labels(db)
        before = project_snapshot(db, project, field_labels=field_labels)
        _apply(db, project, data, user, locked_refs)
        after = project_snapshot(db, project, field_labels=field_labels)
        changed = not same_project_values(before, after)
        if changed:
            audit.record(
                db, user.id, "update", "project", project.id,
                before=before, after=after, author_name=user.full_name,
            )
        if data.financial_entries is not None:
            finances_changed = _replace_financial_entries(db, project, data.financial_entries, references, user)
            changed = changed or finances_changed
        db.commit()
    except IntegrityError:
        _check_integrity_conflict(db, data, project_id)
        raise
    except Exception:
        db.rollback()
        raise
    if changed:
        cache.bump("analytics")
    db.refresh(project)
    return to_out(project)


def _lock_references(
    db: Session, data: ProjectIn | None, entries: list[ProjectFinancialEntryIn],
) -> dict[int, ReferenceValue]:
    if len(entries) > 500:
        raise HTTPException(422, "Не более 500 финансовых строк на проект")
    ids = {
        value for entry in entries for value in (entry.status_id, entry.article_id, entry.cost_type_id)
        if value is not None
    }
    if data is not None:
        ids.update(value for value in (
            data.payment_type_id, data.customer_stage_id, data.indicator_id, data.assessment_status_id,
        ) if value is not None)
    if not ids:
        return {}
    # Все ссылки карточки и финансов блокируются одной транзакцией в одинаковом порядке.
    rows = db.scalars(
        select(ReferenceValue).where(ReferenceValue.id.in_(ids)).order_by(ReferenceValue.id)
        .with_for_update(read=True).execution_options(populate_existing=True)
    ).all()
    return {row.id: row for row in rows}


def _validate_financial_entries(
    entries: list[ProjectFinancialEntryIn],
    locked_refs: dict[int, ReferenceValue],
    project: Project | None = None,
) -> list[tuple[dict[str, ReferenceValue], ProjectFinancialEntry | None]]:
    category_fields = {
        "revenue": {"status_id": "revenue_status"},
        "cost": {"status_id": "cost_status", "article_id": "cost_article", "cost_type_id": "cost_type"},
    }
    matches = []
    old_rows = {row.id: row for row in project.financial_entries} if project else {}
    used_ids: set[int] = set()
    reserved_ids = {entry.id for entry in entries if entry.id is not None}
    fields = ("kind", "year", "month", "amount", "vat_included", "status_id", "article_id", "cost_type_id")
    explicit_ids = ["id" in entry.model_fields_set for entry in entries]
    if any(explicit_ids) and not all(explicit_ids):
        raise HTTPException(422, {"field": "id", "message": (
            "Укажите ID для каждой финансовой строки: число для существующей, null для новой."
        )})
    explicit_protocol = bool(entries) and all(explicit_ids)
    for entry in entries:
        old = old_rows.get(entry.id) if entry.id is not None else None
        if entry.id is not None:
            if old is None or entry.id in used_ids:
                raise HTTPException(422, {"field": "id", "message": "Недопустимый или повторный ID финансовой строки"})
            used_ids.add(entry.id)
        elif not explicit_protocol:
            # Без ID безопасно сохранять идентичность только полностью совпавшей строки.
            old = next((row for row in old_rows.values() if row.id not in used_ids and row.id not in reserved_ids
                        and all(getattr(row, field) == getattr(entry, field) for field in fields)), None)
            if old is not None:
                used_ids.add(old.id)
        if entry.kind == "revenue":
            for field in ("article_id", "cost_type_id"):
                if getattr(entry, field) is not None:
                    raise HTTPException(422, {"field": field, "message": "Поле допустимо только для затрат"})
        row_refs = {}
        for field, category in category_fields[entry.kind].items():
            ref_id = getattr(entry, field)
            if ref_id is None and entry.kind == "cost" and field != "status_id":
                raise HTTPException(422, {"field": field, "message": "Обязательное поле"})
            ref = locked_refs.get(ref_id)
            retained = old is not None and old.kind == entry.kind and getattr(old, field) == ref_id
            if ref is not None and not ref.is_active and not retained and entry.id is None and project is not None:
                raise HTTPException(422, {"field": field, "message": (
                    "Неактивное значение нельзя назначить новой строке. Чтобы изменить существующую строку, "
                    "загрузите карточку заново и передайте её ID."
                )})
            if ref is not None and ((not ref.is_active and not retained) or ref.category != category):
                raise HTTPException(
                    422, {"field": field, "message": "Справочник не соответствует типу финансовой строки"},
                )
            if ref is None:
                raise HTTPException(422, {"field": field, "message": "Значение справочника не найдено"})
            row_refs[field] = ref
        matches.append((row_refs, old))

    # Без ID нельзя отличить редактирование от удаления и добавления, а также
    # выбрать, какую из нескольких одинаковых строк следует удалить.
    unmatched_old = [row for row in old_rows.values() if row.id not in used_ids]
    new_without_id = any(entry.id is None and old is None for entry, (_, old) in zip(entries, matches, strict=True))
    ambiguous_duplicate = any(
        entry.id is None and old is not None
        and any(all(getattr(row, field) == getattr(entry, field) for field in fields) for row in unmatched_old)
        for entry, (_, old) in zip(entries, matches, strict=True)
    )
    if not explicit_protocol and unmatched_old and (new_without_id or ambiguous_duplicate):
        raise HTTPException(422, {"field": "id", "message": (
            "Неоднозначное изменение финансовых строк. Обновите карточку и передайте ID существующих строк; "
            "удаление и добавление без ID выполните раздельно."
        )})
    return matches


def _replace_financial_entries(
    db: Session,
    project: Project,
    entries: list[ProjectFinancialEntryIn],
    references: list[tuple[dict[str, ReferenceValue], ProjectFinancialEntry | None]],
    user: User,
) -> bool:
    # Изменяем сопоставленные строки на месте; новые создаём, исключённые удаляем.
    field_labels = effective_project_field_labels(db)
    before = [financial_snapshot(row, field_labels=field_labels) | {"_row_id": row.id}
              for row in project.financial_entries]
    after = [financial_snapshot(entry, ref, field_labels=field_labels)
             for entry, (ref, _) in zip(entries, references, strict=True)]
    changed = len(project.financial_entries) != len(entries) or any(
        old is None or not same_finances(
            [financial_snapshot(old, field_labels=field_labels)], [snapshot],
        )
        for snapshot, (_, old) in zip(after, references, strict=True)
    )
    if not changed:
        return False
    keep = {old.id for _, old in references if old is not None}
    for old in list(project.financial_entries):
        if old.id not in keep:
            project.financial_entries.remove(old)
    saved_rows = []
    for entry, (ref, old) in zip(entries, references, strict=True):
        row = old if old is not None else ProjectFinancialEntry()
        if old is None:
            project.financial_entries.append(row)
        for field in ("kind", "year", "month", "amount", "vat_included"):
            setattr(row, field, getattr(entry, field))
        row.status = ref["status_id"]
        row.article = ref.get("article_id")
        row.cost_type = ref.get("cost_type_id")
        saved_rows.append(row)
    db.flush()
    after = [snapshot | {"_row_id": row.id} for snapshot, row in zip(after, saved_rows, strict=True)]
    project.updated_at = func.now()
    audit.record(
        db, user.id, "replace_financial_entries", "project", project.id,
        before=before, after=after, author_name=user.full_name,
    )
    return True


@router.put("/{project_id}/financial-entries", response_model=list[ProjectFinancialEntryOut])
def replace_financial_entries(
    project_id: int,
    entries: list[ProjectFinancialEntryIn],
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*PROJECT_EDITORS)),
) -> list[ProjectFinancialEntry]:
    project = _get_visible(db, project_id, user, for_update=True)
    _require_active(project)
    try:
        locked_refs = _lock_references(db, None, entries)
        references = _validate_financial_entries(entries, locked_refs, project)
        changed = _replace_financial_entries(db, project, entries, references, user)
        if not changed:
            return project.financial_entries
        db.commit()
    except Exception:
        db.rollback()
        raise
    cache.bump("analytics")
    db.refresh(project)
    return project.financial_entries
