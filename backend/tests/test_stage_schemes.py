from datetime import date
from io import BytesIO
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook
from project_factories import create_legacy_project, customer_payload, legacy_payload
from sqlalchemy import func, select

from app.api import analytics as analytics_api
from app.api import projects
from app.main import app
from app.models import AuditLog, Project, ReferenceValue, Stage
from app.schemas import LEGACY_STAGE_SCOPE_NOTE
from app.services import analytics
from app.services.project_field_labels import same_project_values
from app.services.project_history import PROJECT_FIELDS, project_snapshot


def _create(client, db, **changes):
    response = client.post("/api/projects", json=customer_payload(db, **changes))
    assert response.status_code == 201, response.text
    return response.json()


def _counts(db):
    return tuple(db.scalar(select(func.count()).select_from(model)) for model in (Project, AuditLog))


def test_new_project_has_only_customer_stage_and_server_probability(as_user, db):
    payload = customer_payload(db, probability=99)
    expected = db.get(ReferenceValue, payload["customer_stage_id"]).probability
    for key in ("stage", "probability"):
        payload.pop(key)
    response = as_user.post("/api/projects", json=payload)
    assert response.status_code == 201, response.text
    project = response.json()
    assert project["stage_scheme"] == "customer"
    assert project["stage"] is None and project["closed_at"] is None
    assert project["probability"] == expected
    assert project["customer_stage"]["id"] == payload["customer_stage_id"]
    assert as_user.get(f"/api/projects/{project['id']}").json() == project


@pytest.mark.parametrize("invalid", ["missing", "null", "unknown", "inactive", "wrong_category", "no_probability"])
def test_customer_create_validates_stage_before_any_write(as_user, db, monkeypatch, invalid):
    payload = customer_payload(db)
    if invalid == "missing":
        payload.pop("customer_stage_id")
    elif invalid == "null":
        payload["customer_stage_id"] = None
    elif invalid == "unknown":
        payload["customer_stage_id"] = -1
    elif invalid == "wrong_category":
        payload["customer_stage_id"] = db.scalar(select(ReferenceValue.id).where(
            ReferenceValue.category == "payment_type",
        ))
    else:
        reference = db.get(ReferenceValue, payload["customer_stage_id"])
        if invalid == "inactive":
            reference.is_active = False
        else:
            reference.probability = None
        db.commit()
    before = _counts(db)
    bump = Mock()
    monkeypatch.setattr(projects.cache, "bump", bump)
    response = as_user.post("/api/projects", json=payload)
    assert response.status_code == 422
    assert response.json()["detail"]["field"] == "customer_stage_id"
    assert _counts(db) == before
    bump.assert_not_called()


@pytest.mark.parametrize("stage", list(Stage))
def test_old_six_stage_post_is_intentionally_rejected(as_user, db, stage):
    before = _counts(db)
    # Наличие явно выбранного customer-stage не разрешает записать прежнюю стадию.
    for payload in (legacy_payload(db, stage=stage.value), customer_payload(db, stage=stage.value)):
        response = as_user.post("/api/projects", json=payload)
        assert response.status_code == 422
        assert response.json()["detail"]["field"] == "stage"
    assert _counts(db) == before


@pytest.mark.parametrize("scheme", ["legacy", "customer", None, "other"])
def test_stage_scheme_is_readonly_for_post_and_both_put_generations(as_user, db, scheme):
    customer = _create(as_user, db)
    legacy = create_legacy_project(db, name="Прежний проект")
    before = _counts(db)
    requests = [
        ("/api/projects", customer_payload(db, stage_scheme=scheme)),
        (f"/api/projects/{customer['id']}", customer_payload(db, stage_scheme=scheme)),
        (f"/api/projects/{legacy.id}", legacy_payload(db, name=legacy.name, stage_scheme=scheme)),
    ]
    for path, body in requests:
        response = (as_user.post if path == "/api/projects" else as_user.put)(path, json=body)
        assert response.status_code == 422
        assert "stage_scheme" in str(response.json()["detail"])
    assert _counts(db) == before
    assert db.get(Project, legacy.id).stage_scheme == "legacy"
    assert db.get(Project, customer["id"]).stage_scheme == "customer"


@pytest.mark.parametrize("stage", list(Stage))
def test_legacy_six_stages_and_probability_remain_editable(as_user, db, stage):
    legacy = create_legacy_project(db, probability=37)
    response = as_user.put(f"/api/projects/{legacy.id}", json=legacy_payload(db, stage=stage.value, probability=37))
    assert response.status_code == 200, response.text
    saved = response.json()
    assert saved["stage_scheme"] == "legacy" and saved["stage"] == stage.value
    assert saved["probability"] == 37
    assert (saved["closed_at"] is not None) == (stage in {Stage.won, Stage.lost})


def test_legacy_auxiliary_customer_stage_and_probability_do_not_change_implicitly(as_user, db):
    stage_id = customer_payload(db)["customer_stage_id"]
    legacy = create_legacy_project(db, customer_stage_id=stage_id, probability=37)
    reference = db.get(ReferenceValue, stage_id)
    reference.probability = 92
    reference.is_active = False
    db.commit()
    for fields in ({}, {"customer_stage_id": stage_id}):
        payload = legacy_payload(db, description="Обычная правка", **fields)
        payload.pop("probability")
        response = as_user.put(f"/api/projects/{legacy.id}", json=payload)
        assert response.status_code == 200, response.text
        assert response.json()["probability"] == 37
        assert response.json()["customer_stage"]["id"] == stage_id
        assert response.json()["stage"] == "proposal"
    before = as_user.get(f"/api/projects/{legacy.id}").json()
    other_id = db.scalar(select(ReferenceValue.id).where(
        ReferenceValue.category == "customer_stage", ReferenceValue.id != stage_id,
    ))
    for forbidden in (None, other_id):
        response = as_user.put(f"/api/projects/{legacy.id}", json=legacy_payload(
            db, description="Не сохранять", customer_stage_id=forbidden,
        ))
        assert response.status_code == 422
        assert response.json()["detail"]["field"] == "customer_stage_id"
        assert as_user.get(f"/api/projects/{legacy.id}").json() == before


def test_customer_probability_changes_only_with_stage_and_noop_has_no_history(as_user, db, monkeypatch):
    saved = _create(as_user, db, probability=99)
    original = saved["probability"]
    reference = db.get(ReferenceValue, saved["customer_stage"]["id"])
    reference.probability = 13
    db.commit()
    before = _counts(db)
    bump = Mock()
    monkeypatch.setattr(projects.cache, "bump", bump)
    response = as_user.put(f"/api/projects/{saved['id']}", json=customer_payload(db, probability=88,
        customer_stage_id=reference.id))
    assert response.status_code == 200, response.text
    assert response.json() == saved
    assert _counts(db) == before
    bump.assert_not_called()
    response = as_user.put(f"/api/projects/{saved['id']}", json=customer_payload(db, probability=99,
        customer_stage_id=reference.id, description="Правка описания"))
    assert response.status_code == 200
    assert response.json()["probability"] == original
    success = db.scalar(select(ReferenceValue).where(
        ReferenceValue.category == "customer_stage", ReferenceValue.code == "success",
    ))
    response = as_user.put(f"/api/projects/{saved['id']}", json=customer_payload(db,
        customer_stage_id=success.id, probability=1))
    assert response.status_code == 200, response.text
    assert response.json()["probability"] == success.probability
    assert response.json()["stage"] is None and response.json()["closed_at"] is None


def test_customer_stage_cannot_be_cleared_or_replaced_by_legacy_stage(as_user, db):
    saved = _create(as_user, db)
    for fields in ({"customer_stage_id": None}, {"stage": "won"}):
        before = _counts(db)
        response = as_user.put(f"/api/projects/{saved['id']}", json=customer_payload(db, **fields))
        assert response.status_code == 422
        assert _counts(db) == before
        assert as_user.get(f"/api/projects/{saved['id']}").json() == saved


def test_registry_contains_both_schemes_but_open_only_is_legacy(as_user, db):
    legacy = create_legacy_project(db, name="Поколения старый")
    customer = _create(as_user, db, name="Поколения новый")
    page = as_user.get("/api/projects", params={"q": "Поколения"}).json()
    assert page["stage_scope"] == "all" and page["scope_note"] is None
    all_rows = page["items"]
    assert {p["id"] for p in all_rows} == {legacy.id, customer["id"]}
    page = as_user.get("/api/projects", params={"q": "Поколения", "open_only": True}).json()
    assert page["stage_scope"] == "legacy" and page["scope_note"] == LEGACY_STAGE_SCOPE_NOTE
    open_rows = page["items"]
    assert [p["id"] for p in open_rows] == [legacy.id]


def test_xlsx_labels_null_legacy_stage_as_not_applicable(as_analyst, db):
    _create(as_analyst, db, name="Новый проект для XLSX")
    response = as_analyst.get("/api/projects/export.xlsx", params={"q": "Новый проект для XLSX"})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.active.cell(2, 8).value == "Не применяется"
    assert workbook.active.cell(2, 17).value is not None
    workbook.close()


def test_snapshot_scheme_is_metadata_and_not_user_history_field(as_user, db):
    saved = _create(as_user, db)
    snapshot = project_snapshot(db, db.get(Project, saved["id"]))
    assert snapshot["_stage_scheme"] == "customer"
    assert "stage_scheme" not in PROJECT_FIELDS and "_stage_scheme" not in PROJECT_FIELDS
    assert "_stage_scheme" not in snapshot["_field_labels"]
    assert same_project_values(snapshot, snapshot | {"_stage_scheme": "legacy"})
    changes = as_user.get(f"/api/projects/{saved['id']}/history").json()["items"][0]["changes"]
    assert not any("scheme" in row["field"] for row in changes)


def test_openapi_stage_scheme_is_readonly_literal_and_stage_nullable():
    schemas = app.openapi()["components"]["schemas"]
    output = schemas["ProjectOut"]["properties"]
    assert output["stage_scheme"]["readOnly"] is True
    assert output["stage_scheme"]["enum"] == ["legacy", "customer"]
    assert {"type": "null"} in output["stage"]["anyOf"]
    assert "stage_scheme" not in schemas["ProjectSaveIn"]["properties"]


def test_all_analytics_ignore_customer_projects_on_cache_miss_and_hit(as_analyst, db, monkeypatch):
    stored = {}
    monkeypatch.setattr(analytics_api.cache, "get_json", stored.get)
    monkeypatch.setattr(analytics_api.cache, "set_json", lambda key, value: stored.__setitem__(key, value))
    paths = ("kpi", "trend", "breakdown")
    before = {path: as_analyst.get(f"/api/analytics/{path}", params={"as_of": "2026-10-21"}).json() for path in paths}
    _create(as_analyst, db, amount="999999999999.99", start_date="2026-01-01")
    for hit in (False, True):
        if not hit:
            stored.clear()
        for path in paths:
            response = as_analyst.get(f"/api/analytics/{path}", params={"as_of": "2026-10-21"})
            assert response.status_code == 200
            assert response.json() == before[path]
            assert response.json()["stage_scope"] == "legacy"
            assert response.json()["scope_note"] == LEGACY_STAGE_SCOPE_NOTE
    frame = analytics.load_frame(db, analytics.Filters())
    assert len(frame) == 300


def test_analytics_cache_keys_include_scope_and_worker_uses_same_key(monkeypatch, db):
    from contextlib import nullcontext

    from app import worker

    saved = {}
    calls = []
    monkeypatch.setattr(worker, "SessionLocal", lambda: nullcontext(db))
    monkeypatch.setattr(worker.cache, "set_json", lambda key, result: saved.__setitem__(key, result))
    for name in ("kpi", "trend", "breakdown"):
        monkeypatch.setattr(analytics, name, lambda *args, **kwargs: analytics.STAGE_SCOPE.copy())
    make_key = analytics.cache.make_key

    def scoped_key(namespace, **parts):
        calls.append(parts)
        return make_key(namespace, **parts)

    monkeypatch.setattr(analytics.cache, "make_key", scoped_key)
    worker.warm_analytics_cache()
    assert len(saved) == 7 and all(parts["stage_scope"] == "legacy" for parts in calls)
    # Подключаем чтение прогретого результата для проверки общего ключа API/Celery.
    monkeypatch.setattr(analytics_api.cache, "get_json", saved.get)
    result = analytics_api.cached("kpi", analytics.Filters(), lambda: pytest.fail("Ключ прогрева не совпал"),
                                  as_of=date.today())
    assert result == analytics.STAGE_SCOPE
