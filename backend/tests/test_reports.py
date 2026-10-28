from datetime import UTC
from decimal import Decimal
from io import BytesIO

import pytest
from conftest import login
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Project, ProjectFinancialEntry, ReferenceValue, Role


def definition(grain="project", columns=None, filters=None):
    return {
        "grain": grain,
        "columns": columns or ["project.id", "project.name"],
        "filters": filters or [],
        "sort": [{"field": "project.id", "direction": "asc"}],
    }


def create(client, spec=None):
    response = client.post("/api/reports", json={"title": "Отчёт", "definition": spec or definition()})
    assert response.status_code == 201, response.text
    return response.json()


def test_registry_and_validation(as_analyst):
    fields = as_analyst.get("/api/report-fields").json()["fields"]
    assert any(f["id"] == "financial.amount" and f["grains"] == ["financial_entry"] for f in fields)
    for spec in [
        definition(columns=["project.id; DROP TABLE users"]),
        definition(columns=["financial.amount"]),
        definition(filters=[{"field": "project.id", "operator": "eq", "value": "1 OR 1=1"}]),
        definition(filters=[{"field": "project.name", "operator": "unknown", "value": "x"}]),
    ]:
        assert as_analyst.post("/api/reports/preview", json={"definition": spec}).status_code == 422


def test_private_public_scope_and_revocation(client, db):
    login(client, db, "analyst@example.com")
    report = create(client)
    rid = report["id"]
    client.post("/api/auth/logout")
    manager = login(client, db, "manager1@example.com")
    assert client.get(f"/api/reports/{rid}").status_code == 404
    assert client.post(f"/api/reports/{rid}/run", json={"version": 1}).status_code == 404
    assert client.post("/api/reports", json={"title": "x", "definition": definition()}).status_code == 403
    client.post("/api/auth/logout")
    login(client, db, "analyst@example.com")
    assert client.post(f"/api/reports/{rid}/publish", json={"published": True, "version": 1}).status_code == 200
    client.post("/api/auth/logout")
    login(client, db, "manager1@example.com")
    run = client.post(f"/api/reports/{rid}/run", json={"version": 2}).json()
    sid = run["snapshot_id"]
    rows = client.get(f"/api/report-snapshots/{sid}").json()["rows"]
    assert rows
    own = set(db.scalars(select(Project.id).where(Project.manager_id == manager.id, Project.archived_at.is_(None))))
    assert {row["project.id"] for row in rows} <= own
    client.post("/api/auth/logout")
    login(client, db, "analyst@example.com")
    client.post(f"/api/reports/{rid}/publish", json={"published": False, "version": 2})
    client.post("/api/auth/logout")
    login(client, db, "manager1@example.com")
    assert client.get(f"/api/report-snapshots/{sid}").status_code == 404
    assert client.get(f"/api/report-snapshots/{sid}/export").status_code == 404


def test_demoted_owner_cannot_read_private_or_old_snapshot(client, db):
    user = login(client, db, "analyst@example.com")
    report = create(client)
    rid = report["id"]
    sid = client.post(f"/api/reports/{rid}/run", json={"version": 1}).json()["snapshot_id"]
    user.role = Role.user
    db.flush()
    assert rid not in [item["id"] for item in client.get("/api/reports").json()["items"]]
    assert client.get(f"/api/reports/{rid}").status_code == 404
    assert client.post(f"/api/reports/{rid}/run", json={"version": 1}).status_code == 404
    assert client.get(f"/api/report-snapshots/{sid}").status_code == 404
    assert client.get(f"/api/report-snapshots/{sid}/export").status_code == 404


def test_other_analyst_cannot_read_or_modify_private_report(client, db):
    from app.models import User

    login(client, db, "analyst@example.com")
    report = create(client)
    client.post("/api/auth/logout")
    other = db.scalar(select(User).where(User.email == "manager2@example.com"))
    other.role = Role.analyst
    db.flush()
    login(client, db, "manager2@example.com")
    rid = report["id"]
    assert client.get(f"/api/reports/{rid}").status_code == 404
    assert client.post(f"/api/reports/{rid}/run", json={"version": 1}).status_code == 404
    assert (
        client.put(f"/api/reports/{rid}", json={"title": "Чужой", "definition": definition(), "version": 1}).status_code
        == 404
    )
    assert client.post(f"/api/reports/{rid}/publish", json={"version": 1, "published": True}).status_code == 404


def test_admin_can_edit_and_duplicate_another_users_report(client, db):
    login(client, db, "analyst@example.com")
    report = create(client)
    client.post("/api/auth/logout")
    admin = login(client, db, "admin@example.com")
    rid = report["id"]
    updated = client.put(f"/api/reports/{rid}", json={"title": "Новый", "definition": definition(), "version": 1})
    assert updated.status_code == 200
    assert updated.json()["owner_id"] == report["owner_id"]
    copy = client.post(f"/api/reports/{rid}/duplicate", json={"title": "Копия"})
    assert copy.status_code == 201
    assert copy.json()["owner_id"] == admin.id
    assert copy.json()["published"] is False


def test_version_conflicts_and_snapshots(as_analyst):
    report = create(as_analyst)
    rid = report["id"]
    snapshot = as_analyst.post(f"/api/reports/{rid}/run", json={"version": 1}).json()
    assert snapshot["report_version"] == 1
    assert snapshot["title"] == "Отчёт"
    payload = {"title": "Новое название", "definition": definition(), "version": 1}
    assert as_analyst.put(f"/api/reports/{rid}", json=payload).json()["version"] == 2
    assert as_analyst.post(f"/api/reports/{rid}/run", json={"version": 1}).status_code == 409
    assert as_analyst.put(f"/api/reports/{rid}", json=payload).status_code == 409
    assert as_analyst.post(f"/api/reports/{rid}/publish", json={"version": 1, "published": True}).status_code == 409
    assert as_analyst.delete(f"/api/reports/{rid}?version=1").status_code == 409
    response = as_analyst.get(f"/api/report-snapshots/{snapshot['snapshot_id']}")
    assert response.json()["report_version"] == 1
    assert response.json()["title"] == "Отчёт"
    assert response.headers["cache-control"] == "private, no-store"
    exported = as_analyst.get(f"/api/report-snapshots/{snapshot['snapshot_id']}/export")
    assert exported.headers["cache-control"] == "private, no-store"
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook.properties.title == "Отчёт"
    assert workbook.active["A1"].value == response.json()["columns"][0]["label"]


def test_snapshot_quota_can_be_released(as_analyst):
    for _ in range(5):
        assert as_analyst.post("/api/reports/preview", json={"definition": definition()}).status_code == 200
    assert as_analyst.post("/api/reports/preview", json={"definition": definition()}).status_code == 429
    assert as_analyst.delete("/api/report-snapshots").status_code == 204
    assert as_analyst.post("/api/reports/preview", json={"definition": definition()}).status_code == 200


def test_archived_project_invalidates_analyst_snapshot(as_analyst, db):
    from datetime import datetime

    project = db.scalars(select(Project).where(Project.archived_at.is_(None))).first()
    spec = definition(filters=[{"field": "project.id", "operator": "eq", "value": project.id}])
    sid = as_analyst.post("/api/reports/preview", json={"definition": spec}).json()["snapshot_id"]
    project.archived_at = datetime.now(UTC)
    db.flush()
    assert as_analyst.get(f"/api/report-snapshots/{sid}").status_code == 404
    assert as_analyst.get(f"/api/report-snapshots/{sid}/export").status_code == 404


def test_published_filter_precedes_catalog_pagination(as_analyst):
    first = create(as_analyst)
    second = create(as_analyst)
    as_analyst.post(f"/api/reports/{second['id']}/publish", json={"published": True, "version": 1})
    page = as_analyst.get("/api/reports", params={"published": "true", "limit": 1}).json()
    assert page["total"] == 1
    assert page["items"][0]["id"] == second["id"]
    assert first["id"] != page["items"][0]["id"]


def test_filter_value_rejects_nested_and_long_values(as_analyst):
    for value in ([1] * 100, {"sql": "x"}, "x" * 501):
        spec = definition(filters=[{"field": "project.name", "operator": "eq", "value": value}])
        assert as_analyst.post("/api/reports/preview", json={"definition": spec}).status_code == 422
    assert as_analyst.post("/api/reports/preview", content=b"x" * 65537).status_code == 413
    spec = definition(filters=[{"field": "project.id", "operator": "eq", "value": 2**53}])
    assert as_analyst.post("/api/reports/preview", json={"definition": spec}).status_code == 422


def test_registry_choices_use_codes_and_labels(as_analyst):
    fields = {item["id"]: item for item in as_analyst.get("/api/report-fields").json()["fields"]}
    assert {"value": "lead", "label": "Лид"} in fields["project.stage"]["choices"]
    assert {"value": "revenue", "label": "Выручка"} in fields["financial.kind"]["choices"]
    spec = definition(filters=[{"field": "project.stage", "operator": "eq", "value": "Лид"}])
    assert as_analyst.post("/api/reports/preview", json={"definition": spec}).status_code == 422
    spec["filters"][0]["value"] = "lead"
    response = as_analyst.post("/api/reports/preview", json={"definition": spec})
    assert response.status_code == 200, response.text


def test_limits_and_expired_snapshot_cleanup(as_analyst, db, monkeypatch):
    from datetime import UTC, datetime, timedelta

    from app.models import ReportSnapshot
    from app.services import report_builder

    monkeypatch.setattr(report_builder, "MAX_ROWS", 2)
    assert as_analyst.post("/api/reports/preview", json={"definition": definition()}).status_code == 422
    monkeypatch.setattr(report_builder, "MAX_ROWS", 10000)
    monkeypatch.setattr(report_builder, "MAX_BYTES", 32)
    assert as_analyst.post("/api/reports/preview", json={"definition": definition()}).status_code == 422
    monkeypatch.setattr(report_builder, "MAX_BYTES", 10 * 1024 * 1024)
    sid = as_analyst.post("/api/reports/preview", json={"definition": definition()}).json()["snapshot_id"]
    db.get(ReportSnapshot, sid).expires_at = datetime.now(UTC) - timedelta(seconds=1)
    db.flush()
    assert as_analyst.get(f"/api/report-snapshots/{sid}").status_code == 404
    report_builder.cleanup(db)
    db.flush()
    assert db.get(ReportSnapshot, sid) is None


def test_renamed_heading_is_frozen_with_snapshot(as_analyst, db):
    from app.models import ProjectFieldLabel

    db.add(ProjectFieldLabel(key="name", label="Особое название"))
    db.flush()
    sid = as_analyst.post("/api/reports/preview", json={"definition": definition()}).json()["snapshot_id"]
    label = db.get(ProjectFieldLabel, "name")
    label.label = "Другая подпись"
    db.flush()
    page = as_analyst.get(f"/api/report-snapshots/{sid}").json()
    assert page["columns"][1]["label"] == "Особое название"
    workbook = load_workbook(BytesIO(as_analyst.get(f"/api/report-snapshots/{sid}/export").content))
    assert workbook.active["B1"].value == "Особое название"


def test_finance_grain_preserves_rows_and_snapshot_export(as_analyst, db):
    project = db.scalars(select(Project).where(Project.archived_at.is_(None))).first()
    revenue = db.scalar(select(ReferenceValue).where(ReferenceValue.category == "revenue_status"))
    cost = db.scalar(select(ReferenceValue).where(ReferenceValue.category == "cost_status"))
    for kind, status, amount in [
        ("revenue", revenue, "12.123"),
        ("revenue", revenue, "15.456"),
        ("cost", cost, "4.001"),
        ("cost", cost, "5.002"),
    ]:
        db.add(
            ProjectFinancialEntry(
                project_id=project.id,
                kind=kind,
                year=2027,
                month=1,
                amount=Decimal(amount),
                vat_included=True,
                status_id=status.id,
            )
        )
    db.flush()
    spec = definition(
        "financial_entry",
        ["project.id", "financial.id", "financial.kind", "financial.amount"],
        [{"field": "project.id", "operator": "eq", "value": project.id}],
    )
    snap = as_analyst.post("/api/reports/preview", json={"definition": spec})
    assert snap.status_code == 200, snap.text
    sid = snap.json()["snapshot_id"]
    data = as_analyst.get(f"/api/report-snapshots/{sid}").json()
    assert {"12.123", "15.456", "4.001", "5.002"} <= {row["financial.amount"] for row in data["rows"]}
    assert len({row["financial.id"] for row in data["rows"]}) == data["total"]
    assert {"Выручка", "Затраты"} <= {row["financial.kind"] for row in data["rows"]}
    workbook = load_workbook(BytesIO(as_analyst.get(f"/api/report-snapshots/{sid}/export").content))
    assert workbook.active.max_row == data["total"] + 1


@pytest.mark.parametrize("bad", ["=1+1", "+cmd", "-cmd", "@cmd", "\t=1+1"])
def test_xlsx_does_not_execute_text(as_analyst, db, bad):
    project = db.scalars(select(Project).where(Project.archived_at.is_(None))).first()
    project.name = bad
    db.flush()
    spec = definition(
        columns=["project.name"], filters=[{"field": "project.id", "operator": "eq", "value": project.id}]
    )
    snap = as_analyst.post("/api/reports/preview", json={"definition": spec}).json()
    response = as_analyst.get(f"/api/report-snapshots/{snap['snapshot_id']}/export")
    cell = load_workbook(BytesIO(response.content)).active["A2"]
    assert cell.data_type != "f"
