import json
from pathlib import Path

import pytest
from conftest import login
from sqlalchemy import func, select
from test_project_card_validation import card_migration  # noqa: F401 — общая изолированная SQLite/PG-фикстура

from app import models
from app.models import AuditLog


@pytest.fixture
def label_client(client):
    return client


def audit_rows(db):
    return db.scalars(select(AuditLog).where(AuditLog.entity == "project_field_label")).all()


def test_defaults_cover_project_history_and_financial_labels():
    from app.project_field_label_catalog import FIELD_LABELS
    from app.schemas import ProjectIn
    from app.services.project_history import PROJECT_FIELDS

    assert set(ProjectIn.model_fields) <= FIELD_LABELS.keys()
    assert set(PROJECT_FIELDS) <= FIELD_LABELS.keys()
    assert {"margin", "margin_pct", "segment", "client_inn", "service_group", "created_at"} <= FIELD_LABELS.keys()
    assert FIELD_LABELS["amount"].suffix == ", ₽"
    assert FIELD_LABELS["financial_entries.amount"].suffix == ", тыс. ₽"
    assert FIELD_LABELS["margin_pct"].suffix == ", %"
    assert FIELD_LABELS["archived_at"].default_label == "Архивирование"


def test_frontend_fallback_is_generated_from_same_catalog():
    from app.project_field_label_catalog import FIELD_LABELS

    path = Path(__file__).parents[2] / "frontend/src/lib/project-field-label-defaults.json"
    assert json.loads(path.read_text(encoding="utf-8")) == {
        key: {"label": spec.default_label, "suffix": spec.suffix} for key, spec in FIELD_LABELS.items()
    }


@pytest.mark.parametrize("method,path,body", [
    ("GET", "", None), ("PUT", "/name", {"label": "Новое название"}), ("DELETE", "/name", None),
])
def test_api_requires_authentication(label_client, method, path, body):
    assert label_client.request(method, f"/api/project-fields{path}", json=body).status_code == 401


@pytest.mark.parametrize("email", ["manager1@example.com", "analyst@example.com", "admin@example.com"])
def test_all_authenticated_roles_read(label_client, db, email):
    login(label_client, db, email)
    response = label_client.get("/api/project-fields")
    assert response.status_code == 200
    labels = {row["key"]: row for row in response.json()}
    assert labels["name"]["label"] == labels["name"]["default_label"]
    assert labels["amount"]["suffix"] == ", ₽"
    assert labels["name"]["is_overridden"] is False


@pytest.mark.parametrize("email", ["manager1@example.com", "analyst@example.com"])
def test_non_admin_cannot_change_or_reset(label_client, db, email):
    login(label_client, db, email)
    assert label_client.put("/api/project-fields/name", json={"label": "Подмена"}).status_code == 403
    assert label_client.delete("/api/project-fields/name").status_code == 403
    assert not audit_rows(db)


def test_admin_persists_trimmed_plain_text_and_audits_only_changes(label_client, db):
    from app.services.project_field_labels import effective_project_field_labels

    actor = login(label_client, db, "admin@example.com")
    raw = '<img src=x onerror="alert(1)">'
    response = label_client.put("/api/project-fields/name", json={"label": f"  {raw}  "})
    assert response.status_code == 200, response.text
    assert response.json()["label"] == raw
    assert response.json()["is_overridden"] is True
    db.expire_all()
    assert effective_project_field_labels(db)["name"] == raw
    rows = audit_rows(db)
    assert len(rows) == 1
    assert rows[0].user_id == actor.id
    assert rows[0].payload["key"] == "name"
    assert rows[0].payload["after"] == raw
    assert label_client.put("/api/project-fields/name", json={"label": raw}).status_code == 200
    assert len(audit_rows(db)) == 1
    response = label_client.delete("/api/project-fields/name")
    assert response.status_code == 200
    assert response.json()["label"] == response.json()["default_label"]
    assert response.json()["is_overridden"] is False
    assert db.scalar(select(func.count()).select_from(models.ProjectFieldLabel)) == 0
    assert len(audit_rows(db)) == 2
    assert label_client.delete("/api/project-fields/name").status_code == 200
    assert len(audit_rows(db)) == 2


def test_saving_default_is_reset_and_initial_default_is_noop(label_client, db):
    from app.project_field_label_catalog import FIELD_LABELS

    login(label_client, db, "admin@example.com")
    default = FIELD_LABELS["name"].default_label
    for value in (default, "Другое имя", default):
        assert label_client.put("/api/project-fields/name", json={"label": value}).status_code == 200
    assert len(audit_rows(db)) == 2
    assert db.scalar(select(func.count()).select_from(models.ProjectFieldLabel)) == 0


@pytest.mark.parametrize("key", ["arbitrary", "password_hash", "_field_labels", "client", "__proto__"])
def test_unknown_keys_are_422(label_client, db, key):
    login(label_client, db, "admin@example.com")
    assert label_client.put(f"/api/project-fields/{key}", json={"label": "Новое имя"}).status_code == 422
    assert label_client.delete(f"/api/project-fields/{key}").status_code == 422
    assert not audit_rows(db)


@pytest.mark.parametrize("body", [
    {}, {"label": ""}, {"label": " \t "}, {"label": "я" * 121}, {"label": None}, {"label": 12},
    {"label": "Метка\nДругая"}, {"label": "Метка\x00"}, {"label": "Скрыто\u202e"},
    {"label": "Допустимо", "key": "amount"}, {"label": "Допустимо", "suffix": "тыс. ₽"},
])
def test_invalid_labels_and_extra_fields_are_rejected(label_client, db, body):
    login(label_client, db, "admin@example.com")
    assert label_client.put("/api/project-fields/name", json=body).status_code == 422
    assert not audit_rows(db)


def test_label_limit_and_spaces(label_client, db):
    login(label_client, db, "admin@example.com")
    for label in ("Я", "я" * 120, "Имя   проекта"):
        response = label_client.put("/api/project-fields/name", json={"label": f"  {label}  "})
        assert response.status_code == 200
        assert response.json()["label"] == label


def test_effective_labels_and_refs_merge_do_not_drop_defaults_or_existing_refs(label_client, db):
    from app.project_field_label_catalog import FIELD_LABELS
    from app.services.project_field_labels import with_project_field_labels

    login(label_client, db, "admin@example.com")
    label_client.put("/api/project-fields/name", json={"label": "Новое имя"})
    original = {"segments": [{"id": 7, "name": "Сегмент"}], "stages": []}
    merged = with_project_field_labels(db, original)
    assert merged["project_field_labels"]["name"] == "Новое имя"
    assert set(merged["project_field_labels"]) == FIELD_LABELS.keys()
    assert merged["project_field_labels"]["amount"] == FIELD_LABELS["amount"].default_label
    assert merged["segments"] == original["segments"]
    assert "project_field_labels" not in original
    response = label_client.get("/api/refs")
    assert response.status_code == 200
    assert response.json()["project_field_labels"] == merged["project_field_labels"]
    assert response.json()["stages"]


def test_snapshot_labels_survive_later_override_and_reset(label_client, db):
    from app.services.project_field_labels import (
        effective_project_field_labels,
        historical_field_label,
        with_snapshot_field_labels,
    )

    login(label_client, db, "admin@example.com")
    label_client.put("/api/project-fields/name", json={"label": "Первое имя"})
    effective = effective_project_field_labels(db)
    source = {"name": "Значение", "_labels": {"client_id": "Клиент тогда"}}
    snapshot = with_snapshot_field_labels(source, effective)
    effective["name"] = "Изменение внешнего словаря"
    label_client.put("/api/project-fields/name", json={"label": "Второе имя"})
    label_client.delete("/api/project-fields/name")
    assert snapshot["name"] == source["name"]
    assert snapshot["_labels"] == source["_labels"]
    assert "_field_labels" not in source
    assert historical_field_label("name", snapshot) == "Первое имя"
    assert historical_field_label("name", source) == "Название проекта"
    assert historical_field_label("description", {}) == "Описание"
    assert historical_field_label("indicator_id", {}) == "Индикатор"
    assert historical_field_label("name", {"_field_labels": {"name": 7}}) == "Название проекта"
    assert historical_field_label("name", {"_field_labels": []}) == "Название проекта"


def test_financial_snapshot_keys_keep_units_and_legacy_names():
    from app.services.project_field_labels import historical_field_label, with_snapshot_field_labels

    snapshot = with_snapshot_field_labels({"kind": "cost", "amount": "123.456"}, {
        "financial_entries.article_id": "Статья на дату события", "unknown": "Игнорировать",
    }, keys=("financial_entries.article_id",))
    assert snapshot["amount"] == "123.456"
    assert historical_field_label("financial_entries.article_id", snapshot) == "Статья на дату события"
    assert historical_field_label("financial_entries.article_id", {}) == "Статья"
    assert "unknown" not in snapshot["_field_labels"]


def test_snapshot_metadata_does_not_make_a_project_edit():
    from app.services.project_field_labels import same_project_values

    before = {"name": "Имя", "amount": "123.00", "_labels": {}, "_field_labels": {"name": "Вчера"}}
    after = {**before, "_labels": {"client_id": "Клиент"}, "_field_labels": {"name": "Сегодня"}}
    assert same_project_values(before, after)
    assert not same_project_values(before, {**after, "name": "Другое имя"})


def test_xlsx_headers_keep_order_units_and_plain_strings():
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    from app.project_field_label_catalog import FIELD_LABELS
    from app.services.project_field_label_export import PROJECT_EXPORT_FIELDS, write_project_export_headers

    labels = {key: spec.default_label for key, spec in FIELD_LABELS.items()}
    labels.update({"name": '=HYPERLINK("https://example.invalid")', "amount": "Договор"})
    workbook = Workbook()
    worksheet = workbook.active
    write_project_export_headers(worksheet, labels)
    assert len(PROJECT_EXPORT_FIELDS) == 30
    assert worksheet.cell(1, PROJECT_EXPORT_FIELDS.index("amount") + 1).value == "Договор, ₽"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    saved = load_workbook(output)
    assert saved.active.cell(1, 2).value == labels["name"]
    assert saved.active.cell(1, 2).data_type == "s"
    assert saved.active.cell(1, 30).value.endswith(" (UTC)")


def test_project_history_uses_recorded_field_labels_after_rename(label_client, db):
    from test_projects import _payload

    login(label_client, db, "admin@example.com")
    label_client.put("/api/project-fields/name", json={"label": "Имя на дату создания"})
    project = label_client.post("/api/projects", json=_payload(db)).json()
    event = db.scalar(select(AuditLog).where(
        AuditLog.entity == "project", AuditLog.entity_id == project["id"], AuditLog.action == "create",
    ))
    assert event.payload["after"]["_field_labels"]["name"] == "Имя на дату создания"
    label_client.put("/api/project-fields/name", json={"label": "Сегодняшнее имя"})
    history = label_client.get(f"/api/projects/{project['id']}/history").json()
    name_change = next(change for change in history["items"][0]["changes"] if change["field"] == "name")
    assert name_change["label"] == "Имя на дату создания"
    label_client.delete("/api/project-fields/name")
    history = label_client.get(f"/api/projects/{project['id']}/history").json()
    saved_name = next(change for change in history["items"][0]["changes"] if change["field"] == "name")
    assert saved_name["label"] == "Имя на дату создания"


def test_legacy_history_keeps_original_labels_even_with_active_override(label_client, db):
    from app.services.project_history import HistoryFormatter

    login(label_client, db, "admin@example.com")
    label_client.put("/api/project-fields/description", json={"label": "Сегодняшняя суть"})
    payload = {"before": {"description": "Было"}, "after": {"description": "Стало"}}
    assert HistoryFormatter(db, []).changes("update", payload)[0].label == "Описание"
    payload["after"]["_field_labels"] = {"description": "Суть на дату события"}
    assert HistoryFormatter(db, []).changes("update", payload)[0].label == "Суть на дату события"


def test_financial_and_archive_events_keep_snapshot_labels(label_client, db):
    from test_projects import _payload

    login(label_client, db, "admin@example.com")
    label_client.put("/api/project-fields/financial_entries.revenue", json={"label": "Поступление тогда"})
    label_client.put("/api/project-fields/archived_at", json={"label": "Архив тогда"})
    project = label_client.post("/api/projects", json=_payload(db)).json()
    refs = label_client.get("/api/refs").json()
    response = label_client.put(f"/api/projects/{project['id']}/financial-entries", json=[{
        "kind": "revenue", "year": 2026, "month": 10, "amount": "123.456", "vat_included": True,
        "status_id": refs["revenue_statuses"][0]["id"],
    }])
    assert response.status_code == 200, response.text
    assert label_client.post(f"/api/projects/{project['id']}/archive").status_code == 200
    label_client.put("/api/project-fields/financial_entries.revenue", json={"label": "Поступление сегодня"})
    label_client.put("/api/project-fields/archived_at", json={"label": "Архив сегодня"})
    history = label_client.get(f"/api/projects/{project['id']}/history").json()["items"]
    archived = next(item for item in history if item["action"] == "archive")
    assert archived["changes"][0]["label"] == "Архив тогда"
    financial = next(item for item in history if item["action"] == "replace_financial_entries")
    assert financial["changes"][0]["after"].startswith("Поступление тогда ·")
    assert "123,456 тыс. ₽" in financial["changes"][0]["after"]
    events = db.scalars(select(AuditLog).where(AuditLog.entity == "project", AuditLog.entity_id == project["id"])).all()
    saved_financial = next(event for event in events if event.action == "replace_financial_entries").payload["after"][0]
    assert saved_financial["_field_labels"]
    assert all(key.startswith("financial_entries.") for key in saved_financial["_field_labels"])
    created = next(event for event in events if event.action == "create").payload["after"]
    assert set(created["_field_labels"]) <= created.keys()
    saved_archive = next(event for event in events if event.action == "archive").payload
    assert set(saved_archive["before"]["_field_labels"]) == {"archived_at"}


def test_labels_migration_upgrade_check_and_downgrade(card_migration):  # noqa: F811 — pytest-фикстура из другого модуля
    from sqlalchemy import MetaData, Table, inspect, text

    from alembic import command

    config, engine, original = card_migration
    command.upgrade(config, "0007")
    command.upgrade(config, "0008")
    command.check(config)
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO project_field_labels (key,label) VALUES ('name','Метка')"))
        assert connection.scalar(text("SELECT label FROM project_field_labels WHERE key='name'")) == "Метка"
        project = Table("projects", MetaData(), autoload_with=connection)
        saved = dict(connection.execute(select(project)).mappings().one())
        assert {key: saved[key] for key in original} == original
    with pytest.raises(RuntimeError, match="подписи"):
        command.downgrade(config, "0007")
    with engine.begin() as connection:
        assert connection.scalar(text("SELECT version_num FROM alembic_version")) == "0008"
        assert connection.scalar(text("SELECT label FROM project_field_labels WHERE key='name'")) == "Метка"
        connection.execute(text("DELETE FROM project_field_labels"))
    command.downgrade(config, "0007")
    assert "project_field_labels" not in inspect(engine).get_table_names()
    command.upgrade(config, "0008")
    command.check(config)
