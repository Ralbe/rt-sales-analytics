import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from conftest import login
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select, text
from test_projects import _payload

from app.models import AuditLog, Project, ProjectFinancialEntry
from app.services.analytics import Filters, load_frame


def _create(client, db):
    response = client.post("/api/projects", json=_payload(db))
    assert response.status_code == 201, response.text
    return response.json()


def _admin(client, db):
    login(client, db, "admin@example.com")


@pytest.mark.parametrize("email", ["analyst@example.com", "manager1@example.com"])
def test_archive_and_restore_are_admin_only(client, db, email):
    login(client, db, email)
    project = _create(client, db)
    for action in ("archive", "restore"):
        assert client.post(f"/api/projects/{project['id']}/{action}").status_code == 403
    assert client.get("/api/projects?archived=true").status_code == 403


def test_archive_preserves_finances_and_history_and_restore_returns_all_data(client, db):
    _admin(client, db)
    project = _create(client, db)
    pid = project["id"]
    refs = client.get("/api/refs").json()
    entry = {"kind": "revenue", "year": 2026, "month": 10, "amount": "123.456",
             "vat_included": True, "status_id": refs["revenue_statuses"][0]["id"]}
    rows = client.put(f"/api/projects/{pid}/financial-entries", json=[entry]).json()
    first = client.post(f"/api/projects/{pid}/archive")
    assert first.status_code == 200, first.text
    assert first.json()["archived_at"]
    assert first.json()["financial_entries"] == rows
    assert db.scalar(select(func.count()).select_from(ProjectFinancialEntry).where(
        ProjectFinancialEntry.project_id == pid)) == 1
    history = client.get(f"/api/projects/{pid}/history").json()["items"]
    assert history[0]["action"] == "archive"
    assert history[0]["title"] == "Архивирование проекта"
    assert any(item["action"] == "create" for item in history)
    restored = client.post(f"/api/projects/{pid}/restore").json()
    assert restored["archived_at"] is None
    assert restored["financial_entries"] == rows
    for field in ("name", "amount", "cost", "probability", "created_at"):
        assert restored[field] == project[field]
    assert client.get(f"/api/projects/{pid}/history").json()["items"][0]["action"] == "restore"


def test_repeated_archive_restore_are_idempotent(client, db):
    _admin(client, db)
    pid = _create(client, db)["id"]
    for action in ("archive", "restore"):
        first = client.post(f"/api/projects/{pid}/{action}")
        second = client.post(f"/api/projects/{pid}/{action}")
        assert first.status_code == second.status_code == 200
        assert first.json()["archived_at"] == second.json()["archived_at"]
        assert db.scalar(select(func.count()).select_from(AuditLog).where(
            AuditLog.entity == "project", AuditLog.entity_id == pid, AuditLog.action == action)) == 1


def test_archived_project_is_hidden_from_non_admin_details_history_and_writes(client, db):
    login(client, db, "manager1@example.com")
    project = _create(client, db)
    pid = project["id"]
    _admin(client, db)
    assert client.post(f"/api/projects/{pid}/archive").status_code == 200
    for email in ("analyst@example.com", "manager1@example.com"):
        login(client, db, email)
        assert client.get(f"/api/projects/{pid}").status_code == 404
        assert client.get(f"/api/projects/{pid}/history").status_code == 404
        assert client.put(f"/api/projects/{pid}", json=_payload(db)).status_code == 404
        assert client.put(f"/api/projects/{pid}/financial-entries", json=[]).status_code == 404


def test_admin_must_restore_before_editing_card_or_finances(client, db):
    _admin(client, db)
    pid = _create(client, db)["id"]
    client.post(f"/api/projects/{pid}/archive")
    assert client.put(f"/api/projects/{pid}", json=_payload(db)).status_code == 409
    assert client.put(f"/api/projects/{pid}/financial-entries", json=[]).status_code == 409
    assert client.get(f"/api/projects/{pid}").status_code == 200
    client.post(f"/api/projects/{pid}/restore")
    assert client.put(f"/api/projects/{pid}", json=_payload(db)).status_code == 200


def test_archive_is_excluded_from_active_register_xlsx_and_analytics(client, db):
    _admin(client, db)
    project = _create(client, db)
    pid = project["id"]
    client.post(f"/api/projects/{pid}/archive")
    query = {"q": project["name"]}
    assert client.get("/api/projects", params=query).json()["items"] == []
    archived = client.get("/api/projects", params=dict(query, archived=True)).json()
    assert [p["id"] for p in archived["items"]] == [pid]
    export = client.get("/api/projects/export.xlsx", params=query)
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    assert workbook.active.max_row == 1
    workbook.close()
    assert pid not in set(load_frame(db, Filters())["id"])
    client.post(f"/api/projects/{pid}/restore")
    assert pid in set(load_frame(db, Filters())["id"])


def test_archived_name_remains_reserved_to_allow_lossless_restore(client, db):
    _admin(client, db)
    project = _create(client, db)
    client.post(f"/api/projects/{project['id']}/archive")
    duplicate = client.post("/api/projects", json=_payload(db))
    assert duplicate.status_code == 409
    assert client.post(f"/api/projects/{project['id']}/restore").status_code == 200
    assert db.get(Project, project["id"]).archived_at is None


def test_unknown_archive_target_returns_404_and_unauthenticated_is_denied(client, db):
    assert client.post("/api/projects/999999/archive").status_code == 401
    _admin(client, db)
    assert client.post("/api/projects/999999/archive").status_code == 404
    assert client.post("/api/projects/999999/restore").status_code == 404


def test_archive_migration_preserves_project_finances_and_guards_unsafe_downgrade(tmp_path):
    url = "sqlite:///" + (tmp_path / "archive-migration.db").as_posix()
    env = dict(os.environ, DATABASE_URL=url, REDIS_URL="")
    backend = Path(__file__).resolve().parents[1]

    def migrate(command, revision):
        return subprocess.run([sys.executable, "-m", "alembic", command, revision], cwd=backend,
                              env=env, capture_output=True, text=True, encoding="utf-8")

    initial = migrate("upgrade", "0006")
    assert initial.returncode == 0, initial.stderr
    engine = create_engine(url)
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO users (id,email,full_name,role,password_hash,totp_enabled,is_active) "
                                "VALUES (101,'migration@example.test','Миграция','admin','test-only',0,1)"))
        connection.execute(text("INSERT INTO segments (id,name) VALUES (101,'Сегмент')"))
        connection.execute(text("INSERT INTO services (id,name) VALUES (101,'Услуга')"))
        connection.execute(text("INSERT INTO clients (id,name,inn,segment_id) VALUES (101,'Клиент','0000000000',101)"))
        connection.execute(text("INSERT INTO projects (id,name,description,client_id,manager_id,service_id,stage,"
                                "amount,cost,probability,start_date,expected_close_date,realized_work,period_plan) "
                                "VALUES (101,'Проект','Сохранить текст',101,101,101,'proposal',1234.56,12.34,50,"
                                "'2026-01-01','2026-12-31','Сделано','Планы')"))
        status = connection.scalar(text("SELECT id FROM reference_values WHERE category='revenue_status' LIMIT 1"))
        connection.execute(text("INSERT INTO project_financial_entries "
                                "(project_id,kind,year,month,amount,vat_included,status_id) "
                                "VALUES (101,'revenue',2026,10,123.456,1,:status)"), {"status": status})
    upgraded = migrate("upgrade", "0007")
    assert upgraded.returncode == 0, upgraded.stderr
    with engine.begin() as connection:
        row = connection.execute(text("SELECT name,amount,archived_at FROM projects WHERE id=101")).one()
        assert row.name == "Проект" and str(row.amount) == "1234.56" and row.archived_at is None
        amount = connection.scalar(text("SELECT amount FROM project_financial_entries WHERE project_id=101"))
        assert str(amount) == "123.456"
        connection.execute(text(
            "UPDATE projects SET archived_at='2026-10-26 10:00:00',archived_by_id=101 WHERE id=101"
        ))
    blocked = migrate("downgrade", "0006")
    assert blocked.returncode != 0
    with engine.begin() as connection:
        assert connection.scalar(text("SELECT version_num FROM alembic_version")) == "0007"
        assert connection.scalar(text("SELECT archived_at FROM projects WHERE id=101")) is not None
        connection.execute(text("UPDATE projects SET archived_at=NULL,archived_by_id=NULL WHERE id=101"))
    downgraded = migrate("downgrade", "0006")
    assert downgraded.returncode == 0, downgraded.stderr
    restored = migrate("upgrade", "0007")
    assert restored.returncode == 0, restored.stderr
    with engine.connect() as connection:
        assert connection.scalar(text("SELECT description FROM projects WHERE id=101")) == "Сохранить текст"
        assert connection.scalar(text("SELECT COUNT(*) FROM project_financial_entries WHERE project_id=101")) == 1
    engine.dispose()
