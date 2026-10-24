import random
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditLog, Client, Project, Service, User
from app.seed import generate_inn


def _payload(db, **overrides):
    client = db.scalars(select(Client).order_by(Client.id)).first()
    service = db.scalars(select(Service).order_by(Service.id)).first()
    data = {
        "name": "Новый тестовый проект",
        "description": "",
        "client_id": client.id,
        "service_id": service.id,
        "stage": "proposal",
        "amount": "1500000.00",
        "cost": "900000.00",
        "probability": 50,
        "start_date": "2026-10-01",
        "expected_close_date": "2026-12-31",
    }
    return data | overrides


def test_project_audit_records_nonfinancial_changes(as_user, db):
    created = as_user.post("/api/projects", json=_payload(db)).json()
    response = as_user.put(
        f"/api/projects/{created['id']}",
        json=_payload(db, description="Обновлённое описание", management_control=True),
    )
    assert response.status_code == 200, response.text
    log = db.scalar(select(AuditLog).where(
        AuditLog.entity == "project", AuditLog.entity_id == created["id"], AuditLog.action == "update",
    ).order_by(AuditLog.id.desc()))
    assert log.user_id == db.scalar(select(User.id).where(User.email == "manager1@example.com"))
    assert log.payload["before"]["description"] == ""
    assert log.payload["after"]["description"] == "Обновлённое описание"
    assert log.payload["before"]["management_control"] is False
    assert log.payload["after"]["management_control"] is True
    assert log.payload["before"]["amount"] == log.payload["after"]["amount"]


def test_analyst_sees_all_and_can_edit(as_analyst, db):
    r = as_analyst.get("/api/projects", params={"size": 5})
    assert r.status_code == 200
    assert r.json()["total"] == 300
    created = as_analyst.post("/api/projects", json=_payload(db))
    assert created.status_code == 201, created.text


def test_user_sees_only_own(as_user, db):
    me = db.scalar(select(User).where(User.email == "manager1@example.com"))
    items = as_user.get("/api/projects", params={"size": 100}).json()["items"]
    assert items and all(p["manager"]["id"] == me.id for p in items)


def test_create_project_and_duplicate(as_user, db):
    r = as_user.post("/api/projects", json=_payload(db))
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["margin"] == "600000.00"
    assert body["margin_pct"] == 40.0
    assert body["closed_at"] is None

    check = as_user.get(
        "/api/projects/check-name", params={"name": "новый  тестовый проект", "client_id": body["client"]["id"]}
    )
    assert check.json() == {"available": False}

    dup = as_user.post("/api/projects", json=_payload(db))
    assert dup.status_code == 409
    assert dup.json()["detail"]["field"] == "name"


def test_validation_errors(as_user, db):
    r = as_user.post("/api/projects", json=_payload(db, expected_close_date="2026-09-01", amount="-1"))
    assert r.status_code == 422
    locs = {tuple(e["loc"]) for e in r.json()["detail"]}
    assert ("body", "amount") in locs


def test_closing_sets_closed_at(as_user, db):
    created = as_user.post("/api/projects", json=_payload(db, name="Проект на закрытие")).json()
    won = as_user.put(f"/api/projects/{created['id']}", json=_payload(db, name="Проект на закрытие", stage="won"))
    assert won.status_code == 200
    assert won.json()["closed_at"] is not None


def test_export_xlsx(as_analyst):
    r = as_analyst.get("/api/projects/export.xlsx", params={"open_only": True})
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(r.content), read_only=True)
    headers = next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True))
    assert "Этап по ТЗ" in headers
    assert "Тип платежа" in headers
    assert "Планы на период" in headers


def test_client_autocomplete_and_inn_validation(as_user, db):
    segment_id = db.scalars(select(Client)).first().segment_id
    inn = generate_inn(random.Random(2026))
    assert as_user.get("/api/clients", params={"q": "7"}).status_code == 200
    bad = as_user.post("/api/clients", json={"name": "ООО Тест", "inn": "1234567890", "segment_id": segment_id})
    assert bad.status_code == 422
    ok = as_user.post("/api/clients", json={"name": "ООО Тест", "inn": inn, "segment_id": segment_id})
    assert ok.status_code == 201, ok.text
    dup = as_user.post("/api/clients", json={"name": "ООО Тест 2", "inn": inn, "segment_id": segment_id})
    assert dup.status_code == 409


def test_user_cannot_access_another_users_project(as_user, db):
    me = db.scalar(select(User).where(User.email == "manager1@example.com"))
    project_id = db.scalar(select(Project.id).where(Project.manager_id != me.id))
    response = as_user.get(f"/api/projects/{project_id}")
    assert response.status_code == 404


def test_user_cannot_export_all_projects(as_user):
    assert as_user.get("/api/projects/export.xlsx").status_code == 403


def test_role_reference_matches_spec(as_user):
    roles = as_user.get("/api/refs").json()["roles"]
    assert {role["id"] for role in roles} == {"admin", "analyst", "user"}


def test_customer_spec_references_and_project_fields(as_user, db):
    refs = as_user.get("/api/refs").json()
    assert [item["probability"] for item in refs["customer_stages"]] == [10, 10, 30, 40, 40, 50, 70, 80, 90, 100]
    assert {item["name"] for item in refs["payment_types"]} == {
        "Инсталляции", "Сервисная", "Оборудование", "Разовые", "Интеграционные проекты",
    }
    assert "Снят" in {item["name"] for item in refs["project_indicators"]}

    selected_stage = next(item for item in refs["customer_stages"] if item["probability"] == 30)
    payment_type = refs["payment_types"][0]
    indicator = next(item for item in refs["project_indicators"] if item["name"] == "Снят")
    assessment = refs["assessment_statuses"][0]
    payload = _payload(
        db,
        name="Проект по требованиям заказчика",
        probability=88,
        customer_stage_id=selected_stage["id"],
        payment_type_id=payment_type["id"],
        indicator_id=indicator["id"],
        accepted_for_forecast=True,
        assessment_status_id=assessment["id"],
        subsidiary_implementation=True,
        management_control=True,
        is_industry_solution=True,
        industry_manager="Иван Петров",
        industry_project_number="IND-42",
        implementation_year=2027,
        project_status="В работе",
        realized_work="Завершено обследование",
        period_plan="Подготовить проект решения",
    )

    response = as_user.post("/api/projects", json=payload)
    assert response.status_code == 201, response.text
    project = response.json()
    assert project["customer_stage"]["id"] == selected_stage["id"]
    assert project["probability"] == selected_stage["probability"]
    assert project["stage"] == "proposal"
    assert project["payment_type"]["id"] == payment_type["id"]
    assert project["indicator"]["name"] == "Снят"
    assert project["assessment_status"]["id"] == assessment["id"]
    assert project["industry_project_number"] == "IND-42"
    assert project["implementation_year"] == 2027
    assert project["realized_work"] == "Завершено обследование"
    assert project["period_plan"] == "Подготовить проект решения"
    assert project["created_at"]


def test_project_conditional_fields_require_parent_flags(as_user, db):
    response = as_user.post(
        "/api/projects",
        json=_payload(db, assessment_status_id=1, industry_manager="Иван Петров"),
    )
    assert response.status_code == 422
    assert response.json()["detail"][0]["loc"] == ["body"]

    response = as_user.post(
        "/api/projects",
        json=_payload(db, industry_manager="Иван Петров"),
    )
    assert response.status_code == 422
    assert response.json()["detail"][0]["loc"] == ["body"]


def test_project_reference_id_must_match_expected_catalog(as_user, db):
    refs = as_user.get("/api/refs").json()
    payment_type_id = refs["payment_types"][0]["id"]
    response = as_user.post(
        "/api/projects",
        json=_payload(db, customer_stage_id=payment_type_id),
    )
    assert response.status_code == 422
    assert response.json()["detail"]["field"] == "customer_stage_id"
